# -*- coding: utf-8 -*-
"""Comparison workbook (RTL Arabic). Usage: build_xlsx.py <base> -> <base>/مقارنة_المنتجات.xlsx"""
import json, glob, os, sys
sys.path.insert(0, os.path.dirname(__file__))
from _common import base_dir, data_dir, today, iso_days
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def names_map(base):
    meta = os.path.join(data_dir(base), "study_meta.json")
    if os.path.exists(meta):
        return {int(k): v for k, v in json.load(open(meta)).get("names_ar", {}).items()}
    # fallback: derive from product folder slugs
    m = {}
    for d in sorted(glob.glob(os.path.join(base, "products", "product-*"))):
        b = os.path.basename(d)
        try:
            m[int(b.split("-")[1])] = b.split("-", 2)[2].replace("-", " ")
        except Exception:
            pass
    return m

def prices(base):
    p = os.path.join(data_dir(base), "pricing_research.json")
    out = {}
    if os.path.exists(p):
        for k, v in json.load(open(p)).get("product_prices_lyd", {}).items():
            pn = k.split("_")[0]
            f = v.get("findings") if isinstance(v, dict) else None
            out[pn] = [x["price_lyd"] for x in f if isinstance(x, dict) and x.get("price_lyd")] if isinstance(f, list) else []
    return out

def metrics(base):
    m = {}; TD = today()
    for path in sorted(glob.glob(os.path.join(data_dir(base), "hunt_P*.json"))):
        d = json.load(open(path)); prod = d.get("product") or os.path.basename(path)[5:8]
        comps = d.get("competitors", []); rel = [c for c in comps if c.get("classification") == "relevant"]
        total_ads = sum(len(c.get("ads", []) or []) for c in comps)
        oldest = None; longrun = 0
        for c in comps:
            for ad in c.get("ads", []) or []:
                dd = iso_days(ad.get("start", ""), TD)
                if dd is None: continue
                oldest = dd if oldest is None else max(oldest, dd)
                if dd > 60: longrun += 1
        m[prod] = dict(all=len(comps), rel=len(rel), ads=total_ads, maxdur=oldest or 0, longrun=longrun)
    return m

def main():
    base = base_dir(); DATA = data_dir(base)
    scores = json.load(open(os.path.join(DATA, "scores.json")))
    names = names_map(base); pr = prices(base); met = metrics(base)
    order = sorted(scores.keys())

    wb = Workbook(); ws = wb.active; ws.title = "مقارنة المنتجات"; ws.sheet_view.rightToLeft = True
    headers = ["المنتج", "الرقم", "عدد المنافسين (كل)", "عدد المنافسين المباشرين", "إجمالي الإعلانات",
               "أطول مدة نشاط (يوم)", "إعلانات >60 يوم", "سعر ليبي مرصود (د.ل)", "تكلفة+شحن",
               "الهامش/طلب مُسلَّم", "أعلى خطر", "الحكم /10", "ملاحظة"]
    hf = PatternFill("solid", fgColor="C05621"); hfont = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D0D0D0"); border = Border(thin, thin, thin, thin)
    wrap = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for j, h in enumerate(headers, 1):
        c = ws.cell(1, j, h); c.fill = hf; c.font = hfont; c.alignment = wrap; c.border = border
    for i, pnn in enumerate(order, start=2):
        num = int(pnn[1:]); sc = scores[pnn]; mm = met.get(pnn, {}); v = sc.get("verdict")
        pl = pr.get(pnn, []); price = "، ".join(str(x) for x in pl) if pl else "غير متوفر"
        row = [names.get(num, pnn), pnn, mm.get("all", 0), mm.get("rel", 0), mm.get("ads", 0),
               mm.get("maxdur", 0), mm.get("longrun", 0), price, sc.get("cost_ship", "غير متوفر"),
               sc.get("margin", "غير متوفر"), sc.get("risk_top", "غير متوفر"),
               v if v is not None else "غير متوفر", sc.get("note", "")]
        for j, val in enumerate(row, 1):
            c = ws.cell(i, j, val); c.alignment = wrap; c.border = border
        if isinstance(v, (int, float)):
            col = "C6EFCE" if v >= 7 else ("FFEB9C" if v >= 5 else "FFC7CE")
            ws.cell(i, 12).fill = PatternFill("solid", fgColor=col)
    for j, w in enumerate([26, 6, 13, 16, 14, 15, 13, 18, 16, 20, 20, 9, 38], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 44; ws.freeze_panes = "A2"
    out = os.path.join(base, "مقارنة_المنتجات.xlsx"); wb.save(out)
    print("wrote", out)

if __name__ == "__main__":
    main()
